"""
Resume Parser - Main parsing logic.

This module contains the core resume parsing functionality
to extract text from various file formats.
"""

from pathlib import Path
from typing import Dict, Optional
import logging
import fitz  # PyMuPDF
import docx
import pdfplumber

logger = logging.getLogger(__name__)


class ResumeParser:
    """
    Main resume parser class.

    Handles parsing of resume files in various formats (PDF, DOCX)
    and extracts raw text content.
    """

    def __init__(self):
        """Initialize the resume parser."""
        self.supported_formats = ['.pdf', '.docx', '.doc']
        logger.info("ResumeParser initialized")

    def parse(self, file_path: Path) -> Dict[str, str]:
        """
        Parse a resume file and extract text content.

        Args:
            file_path: Path to the resume file

        Returns:
            Dictionary containing parsed resume data

        Raises:
            ValueError: If file format is not supported
            FileNotFoundError: If file does not exist
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_extension = file_path.suffix.lower()
        if file_extension not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {file_extension}")

        logger.info(f"Parsing resume: {file_path.name}")

        try:
            # Extract text based on file type
            if file_extension == '.pdf':
                raw_text = self.parse_pdf(file_path)
            elif file_extension in ['.docx', '.doc']:
                raw_text = self.parse_docx(file_path)
            else:
                raw_text = ""

            return {
                "file_name": file_path.name,
                "raw_text": raw_text,
                "status": "success",
                "text_length": len(raw_text)
            }

        except Exception as e:
            logger.error(f"Error parsing resume {file_path.name}: {e}")
            return {
                "file_name": file_path.name,
                "raw_text": "",
                "status": "error",
                "error": str(e)
            }

    def parse_pdf(self, file_path: Path) -> str:
        """
        Parse PDF resume file using PyMuPDF and pdfplumber.

        Args:
            file_path: Path to PDF file

        Returns:
            Extracted text content
        """
        text = ""

        try:
            # Try PyMuPDF first (faster)
            doc = fitz.open(file_path)
            for page in doc:
                text += page.get_text()
            doc.close()

            # If PyMuPDF didn't extract much text, try pdfplumber
            if len(text.strip()) < 100:
                logger.info(f"PyMuPDF extracted little text, trying pdfplumber for {file_path.name}")
                text = ""
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"

            logger.info(f"Extracted {len(text)} characters from PDF: {file_path.name}")
            return text.strip()

        except Exception as e:
            logger.error(f"Error parsing PDF {file_path.name}: {e}")
            raise

    def parse_docx(self, file_path: Path) -> str:
        """
        Parse DOCX resume file using python-docx.

        Args:
            file_path: Path to DOCX file

        Returns:
            Extracted text content
        """
        try:
            doc = docx.Document(file_path)
            text = ""

            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                text += "\n"

            logger.info(f"Extracted {len(text)} characters from DOCX: {file_path.name}")
            return text.strip()

        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path.name}: {e}")
            raise
